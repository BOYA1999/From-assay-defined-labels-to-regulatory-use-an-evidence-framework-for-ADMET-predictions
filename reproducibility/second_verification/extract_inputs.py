from pathlib import Path
import csv
import hashlib
import json
import sys

import numpy as np
import openpyxl
import xlrd

base = Path(__file__).resolve().parent
originals = Path(sys.argv[1])
out = Path.cwd() / 'extracted_inputs'
out.mkdir(exist_ok=True)
manifest = json.loads((base / 'source_manifest.json').read_text(encoding='utf-8'))
for name, record in manifest['sources'].items():
    assert hashlib.sha256((originals / name).read_bytes()).hexdigest() == record['sha256'], name

def write(name, fields, rows):
    with (out / name).open('w', encoding='utf-8', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)

with (originals / 'HepClear.csv').open(encoding='utf-8-sig', newline='') as handle:
    records = list(csv.DictReader(handle))
indices = np.arange(len(records))
np.random.RandomState(0).shuffle(indices)
folds = np.full(10, len(records) // 10)
folds[:len(records) % 10] += 1
assignment = {}
start = 0
for fold, size in enumerate(folds, 1):
    assignment.update({int(i): fold for i in indices[start:start + size]})
    start += size
write('deepdelta_identity_folds.csv', ['data_row', 'raw_smiles_sha256', 'seed0_test_fold'], [
    {'data_row': i + 1, 'raw_smiles_sha256': hashlib.sha256(row['SMILES'].encode('utf-8')).hexdigest(),
     'seed0_test_fold': assignment[i]} for i, row in enumerate(records)])

book = xlrd.open_workbook(str(originals / 'ref22_mmc2.xls'))
ppb, summaries = [], []
for name, n, observed, first_prediction in [('25 new measurements_set', 25, 2, 3), ('PFC_test_set', 63, 5, 7), ('PFC99_test_set', 36, 5, 7)]:
    sheet = book.sheet_by_name(name)
    for col in range(first_prediction, first_prediction + 14, 2):
        model = sheet.cell_value(0, col)
        for row in range(2, n + 2):
            ppb.append({'dataset': name, 'model': model, 'source_row': row + 1,
                        'observed_cell': f'{openpyxl.utils.get_column_letter(observed + 1)}{row + 1}',
                        'predicted_cell': f'{openpyxl.utils.get_column_letter(col + 1)}{row + 1}',
                        'observed_ppb_percent': sheet.cell_value(row, observed),
                        'predicted_ppb_percent': sheet.cell_value(row, col)})
        summaries.append({'dataset': name, 'model': model, 'source_cell': f'{openpyxl.utils.get_column_letter(col + 2)}{n + 3}',
                          'source_labelled_r2': sheet.cell_value(n + 2, col + 1)})
write('ppb_observed_predicted.csv', list(ppb[0]), ppb)
write('ppb_reported_summaries.csv', list(summaries[0]), summaries)

sheet = openpyxl.load_workbook(originals / 'ref58_source.xlsx', data_only=True)['Figure 5d, 5e']
drugs = ['Cimetidine', 'Digoxin', 'Erythromycin', 'Methotrexate', 'Nafcillin', 'SN-38', 'Vincristine']
bile = []
for donor, first in [('JFC', 3), ('HUM181001B', 27)]:
    for i, drug in enumerate(drugs):
        row = first + i * 3
        bile.append({'donor': donor, 'drug': drug, 'replicate_cells': f'E{row}:E{row + 2}',
                     'replicate_1': sheet[f'E{row}'].value, 'replicate_2': sheet[f'E{row + 1}'].value,
                     'replicate_3': sheet[f'E{row + 2}'].value, 'cached_mean': sheet[f'F{row}'].value,
                     'cached_sd': sheet[f'G{row}'].value, 'human_reference_cell': f'J{i + 3}',
                     'human_reference': sheet[f'J{i + 3}'].value})
write('biliary_seven_drug_inputs.csv', list(bile[0]), bile)

sheet = openpyxl.load_workbook(originals / 'R42_MOESM5.xlsx', data_only=False)['Fig 4']
definitions = [('Solubility', 'I', [('Deep-PK', 'L'), ('ADMET-AI', 'O'), ('ADMETlab 3.0', 'R')]),
               ('Plasma T1/2', 'U', [('ADMETlab 3.0', 'AD')]),
               ('Permeability', 'AG', [('Deep-PK', 'AJ'), ('ADMET-AI', 'AM'), ('ADMETlab 3.0', 'AP')]),
               ('Microsomes T1/2', 'AS', [('Deep-PK', 'AV'), ('ADMET-AI', 'AY'), ('ADMETlab 3.0', 'BB')]),
               ('hERG', 'BE', [('Deep-PK', 'BH'), ('ADMET-AI', 'BK'), ('ADMETlab 3.0', 'BN')]),
               ('Fraction unbound', 'BQ', [('Deep-PK', 'BT'), ('ADMET-AI', 'BW'), ('ADMETlab 3.0', 'BZ')])]
columns = ['A', 'B'] + sorted({c for _, x, models in definitions for c in [x] + [y for _, y in models]}, key=openpyxl.utils.column_index_from_string)
write('nature_supplied_columns.csv', ['source_row'] + columns,
      [dict(source_row=row, **{c: sheet[f'{c}{row}'].value for c in columns}) for row in range(229, 486)])
reported = {(sheet[f'B{r}'].value, sheet[f'A{r}'].value): (r, sheet[f'C{r}'].value) for r in range(492, 508)}
mapping = [{'endpoint': endpoint, 'model': model, 'experimental_column': x, 'prediction_column': y,
            'reported_rho_cell': f'C{reported[endpoint, model][0]}', 'source_reported_rho': reported[endpoint, model][1]}
           for endpoint, x, models in definitions for model, y in models]
write('nature_field_mapping.csv', list(mapping[0]), mapping)
print(json.dumps({'files': {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(out.glob('*.csv'))}}, indent=2))
